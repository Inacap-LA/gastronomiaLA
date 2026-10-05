import io
import json
import logging
import os
from contextlib import asynccontextmanager
from datetime import datetime, date
from typing import List, Optional, Any, Dict

import firebase_admin
from firebase_admin import credentials, firestore
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, PatternFill
from pydantic import BaseModel, ConfigDict
from reportlab.lib import colors
from reportlab.lib.pagesizes import letter
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle
from sqlalchemy.exc import SQLAlchemyError
from sqlalchemy.orm import Session

from fastapi import Depends, FastAPI, HTTPException, Query, Request, status
from fastapi.responses import HTMLResponse, RedirectResponse, StreamingResponse
from fastapi.staticfiles import StaticFiles
from fastapi.templating import Jinja2Templates

from database import engine, get_db
import models

# ==========================================
# CONFIGURACIÓN DE LOGGING Y DIRECTORIOS
# ==========================================
logging.basicConfig(level=logging.INFO)
logger = logging.getLogger("gastronomia")

os.makedirs("static", exist_ok=True)


# ==========================================
# INICIALIZACIÓN DE FIREBASE FIRESTORE
# ==========================================
db_firestore = None
try:
    if not firebase_admin._apps:
        firebase_json_str = os.getenv("FIREBASE_CREDENTIALS")
        cred_path = os.getenv("FIREBASE_CREDENTIALS_PATH")

        if firebase_json_str:
            # Entornos como Render (JSON inyectado por Variable de Entorno)
            cred_dict = json.loads(firebase_json_str)
            cred = credentials.Certificate(cred_dict)
            firebase_admin.initialize_app(cred)
            logger.info("Firebase inicializado desde variable de entorno FIREBASE_CREDENTIALS.")
        elif cred_path and os.path.exists(cred_path):
            # Entornos con archivo local / PythonAnywhere (credentials.json)
            cred = credentials.Certificate(cred_path)
            firebase_admin.initialize_app(cred)
            logger.info(f"Firebase inicializado desde archivo: {cred_path}")
        else:
            # Desarrollo Local (gcloud Application Default Credentials)
            firebase_admin.initialize_app()
            logger.info("Firebase inicializado con credenciales por defecto de gcloud (ADC).")
            
    db_firestore = firestore.client()
except Exception as e:
    logger.warning(f"No se pudo inicializar Firebase Admin SDK (modo local o fallback activo): {e}")


# ==========================================
# CICLO DE VIDA Y FASTAPI
# ==========================================
@asynccontextmanager
async def lifespan(app: FastAPI):
    """Manejador del ciclo de vida para verificar/crear tablas al arrancar."""
    if engine:
        try:
            models.Base.metadata.create_all(bind=engine)
            logger.info("Tablas de la BD verificadas/creadas exitosamente.")
        except Exception as e:
            logger.error(f"Error al inicializar la base de datos en el arranque: {e}")
    yield


app = FastAPI(
    title="Gestión de Aseo Gastronomía",
    description=(
        "Sistema para registro de checklist de aseo, control de talleres, "
        "gestión de borrador/pendientes de revisión, aprobación por Pañol y exportación de reportes."
    ),
    version="1.8.5",
    lifespan=lifespan,
)

app.mount("/static", StaticFiles(directory="static"), name="static")
templates = Jinja2Templates(directory="templates")


# ==========================================
# ADAPTADOR WSGI PARA PYTHONANYWHERE
# ==========================================
try:
    from a2wsgi import ASGItoWSGI
    wsgi_app = ASGItoWSGI(app)
    logger.info("Adaptador WSGI (a2wsgi) cargado correctamente para soporte de PythonAnywhere.")
except ImportError:
    wsgi_app = None
    logger.info("a2wsgi no está instalado; ejecutando en modo nativo ASGI.")


# ==========================================
# FUNCIONES DE AYUDA (HELPERS)
# ==========================================
def parse_date(date_str: Optional[str]) -> Optional[date]:
    """Convierte un string YYYY-MM-DD a un objeto date de manera segura."""
    if not date_str:
        return None
    try:
        return datetime.strptime(date_str.strip(), "%Y-%m-%d").date()
    except ValueError:
        return None


def safe_json_loads(json_data: Any) -> List[Dict]:
    """Carga de manera segura estructuras o cadenas JSON."""
    if not json_data:
        return []
    if isinstance(json_data, list):
        return json_data
    try:
        return json.loads(json_data)
    except (json.JSONDecodeError, TypeError):
        return []


def obtener_plantillas(db: Session) -> List[Dict]:
    """Obtiene las plantillas desde SQL o cae a Firestore como fallback."""
    try:
        plantillas_db = db.query(models.Plantilla).all()
        if plantillas_db:
            return [
                {
                    "id": p.id,
                    "nombre": p.nombre,
                    "actividades": [a.descripcion for a in p.actividades],
                }
                for p in plantillas_db
            ]
    except SQLAlchemyError as e:
        logger.error(f"Error al consultar plantillas en SQL: {e}")

    # Fallback a Firestore
    if db_firestore:
        try:
            docs = db_firestore.collection("plantillas").stream()
            plantillas_fs = []
            for doc in docs:
                data = doc.to_dict()
                plantillas_fs.append({
                    "id": doc.id,
                    "nombre": data.get("nombre", doc.id),
                    "actividades": data.get("actividades", []),
                })
            if plantillas_fs:
                return plantillas_fs
        except Exception as e:
            logger.error(f"Error al consultar plantillas desde Firestore: {e}")

    return []


# ==========================================
# ESQUEMAS PYDANTIC (V2)
# ==========================================
class DocenteOut(BaseModel):
    rut: str
    nombre: str

    model_config = ConfigDict(from_attributes=True)


class SeccionOut(BaseModel):
    cod_asignatura: Optional[str] = None
    asignatura: str
    seccion: str

    model_config = ConfigDict(from_attributes=True)


class AlumnoOut(BaseModel):
    rut: str
    nombre: str

    model_config = ConfigDict(from_attributes=True)


# ==========================================
# RUTAS DE VISTAS WEB (DOCENTE & INICIO)
# ==========================================

@app.get("/", response_class=HTMLResponse, tags=["Vistas Web"])
async def index(request: Request, db: Session = Depends(get_db)):
    """Página principal: Muestra plantillas activas e historial de revisiones."""
    plantillas = obtener_plantillas(db)

    checklists_db = (
        db.query(models.Checklist)
        .order_by(models.Checklist.fecha_creacion.desc())
        .limit(15)
        .all()
    )

    revisiones = []
    for c in checklists_db:
        actividades = safe_json_loads(getattr(c, "actividades_json", None))
        revisiones.append({
            "id": c.id,
            "taller_id": c.taller,
            "docente": c.rut_docente or "N/A",
            "clase": c.codigo_seccion or "N/A",
            "fecha": c.fecha.strftime("%Y-%m-%d") if c.fecha else "N/A",
            "encargado_taller": getattr(c, "encargado_taller", "N/A") or "N/A",
            "nombre_panolero": getattr(c, "nombre_panolero", "N/A") or "N/A",
            "estado": c.estado,
            "actividades": actividades,
            "fecha_actualizacion": (
                c.fecha_actualizacion.strftime("%Y-%m-%d %H:%M")
                if getattr(c, "fecha_actualizacion", None)
                else ""
            ),
        })

    return templates.TemplateResponse(
        request=request,
        name="index.html",
        context={"plantillas": plantillas, "revisiones": revisiones},
    )


@app.get("/checklist/{taller_id}", response_class=HTMLResponse, tags=["Vistas Web"])
async def nuevo_checklist(request: Request, taller_id: str, db: Session = Depends(get_db)):
    """Vista de formulario para crear un nuevo checklist."""
    plantillas = obtener_plantillas(db)
    plantilla = next((p for p in plantillas if str(p["id"]) == str(taller_id)), None)

    if not plantilla:
        return RedirectResponse(url="/", status_code=status.HTTP_303_SEE_OTHER)

    fecha_actual = datetime.now().strftime("%Y-%m-%d")

    return templates.TemplateResponse(
        request=request,
        name="checklist_form.html",
        context={
            "taller": plantilla,
            "checklist": None,
            "checklist_json": None,
            "fecha_actual": fecha_actual,
        },
    )


@app.get("/checklist/editar/{checklist_id}", response_class=HTMLResponse, tags=["Vistas Web"])
async def editar_checklist(request: Request, checklist_id: int, db: Session = Depends(get_db)):
    """Vista para reabrir y continuar un borrador."""
    checklist = db.query(models.Checklist).filter(models.Checklist.id == checklist_id).first()
    if not checklist:
        raise HTTPException(status_code=404, detail="Checklist no encontrado")

    plantillas = obtener_plantillas(db)
    plantilla = next((p for p in plantillas if str(p["id"]) == str(checklist.taller)), None)
    
    taller_nombre = plantilla["nombre"] if plantilla else f"Taller {checklist.taller}"
    actividades_base = plantilla["actividades"] if plantilla else []

    taller_data = {
        "id": checklist.taller,
        "nombre": taller_nombre,
        "actividades": actividades_base,
    }

    actividades_guardadas = safe_json_loads(getattr(checklist, "actividades_json", None))

    checklist_dict = {
        "id": checklist.id,
        "taller_id": checklist.taller,
        "rut_docente": checklist.rut_docente,
        "codigo_seccion": checklist.codigo_seccion,
        "fecha": checklist.fecha.strftime("%Y-%m-%d") if checklist.fecha else "",
        "encargado_taller": getattr(checklist, "encargado_taller", ""),
        "estado": checklist.estado,
        "actividades": actividades_guardadas,
    }

    return templates.TemplateResponse(
        request=request,
        name="checklist_form.html",
        context={
            "taller": taller_data,
            "checklist": checklist,
            "checklist_json": json.dumps(checklist_dict, ensure_ascii=False),
            "fecha_actual": (
                checklist.fecha.strftime("%Y-%m-%d")
                if checklist.fecha
                else datetime.now().strftime("%Y-%m-%d")
            ),
        },
    )


@app.post("/guardar-checklist", tags=["Vistas Web"])
async def guardar_checklist(request: Request, db: Session = Depends(get_db)):
    """Procesa el guardado como Borrador o el envío para Revisión de Pañol."""
    form_data = await request.form()

    raw_checklist_id = form_data.get("checklist_id")
    taller_id = form_data.get("taller_id")
    docente = form_data.get("docente", "")
    clase = form_data.get("clase", "")
    fecha_str = form_data.get("fecha")
    encargado_taller = form_data.get("encargado_taller", "")
    accion = form_data.get("accion", "borrador")

    plantillas = obtener_plantillas(db)
    plantilla = next((p for p in plantillas if str(p["id"]) == str(taller_id)), None)
    if not plantilla:
        return RedirectResponse(url="/", status_code=status.HTTP_303_SEE_OTHER)

    actividades_raw = plantilla["actividades"]
    actividades_evaluadas = []

    for idx, act in enumerate(actividades_raw):
        actividades_evaluadas.append({
            "actividad": act,
            "alumno_encargado": form_data.get(f"alumno_{idx}", ""),
            "estado": form_data.get(f"estado_{idx}", "CONFORME"),
            "observaciones": form_data.get(f"obs_{idx}", ""),
        })

    rut_docente = docente.split("(")[-1].replace(")", "").strip() if docente and "(" in docente else docente
    codigo_seccion = clase.split(" - ")[0].strip() if clase and " - " in clase else clase
    fecha_obj = parse_date(fecha_str)

    if accion == "enviar_panol":
        nuevo_estado = "PENDIENTE_REVISION"
    elif accion == "finalizar":
        nuevo_estado = "FINALIZADO"
    else:
        nuevo_estado = "BORRADOR"

    checklist = None
    if raw_checklist_id and raw_checklist_id.strip():
        try:
            cid_int = int(raw_checklist_id.strip())
            checklist = db.query(models.Checklist).filter(models.Checklist.id == cid_int).first()
        except ValueError:
            pass

    if not checklist:
        checklist = models.Checklist()
        db.add(checklist)

    checklist.taller = taller_id
    checklist.rut_docente = rut_docente or docente
    checklist.codigo_seccion = codigo_seccion or clase
    checklist.fecha = fecha_obj
    checklist.estado = nuevo_estado

    if hasattr(checklist, "encargado_taller"):
        checklist.encargado_taller = encargado_taller
    if hasattr(checklist, "actividades_json"):
        checklist.actividades_json = json.dumps(actividades_evaluadas, ensure_ascii=False)

    try:
        db.commit()
        db.refresh(checklist)
    except SQLAlchemyError as e:
        db.rollback()
        logger.error(f"Error en BD al guardar checklist: {e}")
        raise HTTPException(
            status_code=status.HTTP_500_INTERNAL_SERVER_ERROR,
            detail="Error de base de datos al guardar el checklist.",
        )

    if nuevo_estado == "PENDIENTE_REVISION":
        return RedirectResponse(url="/?msg=enviado_a_panol", status_code=status.HTTP_303_SEE_OTHER)
    elif nuevo_estado == "FINALIZADO":
        return RedirectResponse(url="/?msg=finalizado", status_code=status.HTTP_303_SEE_OTHER)
    else:
        return RedirectResponse(
            url=f"/checklist/editar/{checklist.id}?msg=borrador_guardado",
            status_code=status.HTTP_303_SEE_OTHER,
        )


# ==========================================
# RUTAS DE REVISIÓN PAÑOLERO (DOBLE CHECKLIST)
# ==========================================

@app.get("/checklist/revisar-panol/{checklist_id}", response_class=HTMLResponse, tags=["Pañol"])
async def vista_revision_panol(checklist_id: int, request: Request, db: Session = Depends(get_db)):
    """Vista para que el Pañolero revise y recepcione el taller con Doble Checklist."""
    checklist = db.query(models.Checklist).filter(models.Checklist.id == checklist_id).first()
    if not checklist:
        raise HTTPException(status_code=404, detail="Checklist no encontrado")

    actividades = safe_json_loads(getattr(checklist, "actividades_json", None))

    return templates.TemplateResponse(
        request=request,
        name="revision_panol.html",
        context={"checklist": checklist, "actividades": actividades}
    )


@app.post("/checklist/aprobar-panol/{checklist_id}", tags=["Pañol"])
async def aprobar_revision_panol(
    checklist_id: int,
    request: Request,
    db: Session = Depends(get_db)
):
    """Registra la revisión del Pañolero (Doble Checklist) y finaliza el registro."""
    form_data = await request.form()
    nombre_panolero = form_data.get("nombre_panolero", "")
    observacion_panolero = form_data.get("observacion_panolero", "")

    checklist = db.query(models.Checklist).filter(models.Checklist.id == checklist_id).first()
    if not checklist:
        raise HTTPException(status_code=404, detail="Checklist no encontrado")

    actividades_docente = safe_json_loads(getattr(checklist, "actividades_json", None))
    actividades_actualizadas = []

    for idx, item in enumerate(actividades_docente):
        estado_panol = form_data.get(f"estado_panol_{idx}", "CONFORME")
        obs_panol = form_data.get(f"obs_panol_{idx}", "")

        item["estado_panol"] = estado_panol
        item["obs_panol"] = obs_panol
        actividades_actualizadas.append(item)

    if hasattr(checklist, "nombre_panolero"):
        checklist.nombre_panolero = nombre_panolero
    if hasattr(checklist, "observacion_panolero"):
        checklist.observacion_panolero = observacion_panolero
    if hasattr(checklist, "fecha_revision_panolero"):
        checklist.fecha_revision_panolero = datetime.now()
    if hasattr(checklist, "actividades_json"):
        checklist.actividades_json = json.dumps(actividades_actualizadas, ensure_ascii=False)

    checklist.estado = "FINALIZADO"

    try:
        db.commit()
    except SQLAlchemyError as e:
        db.rollback()
        logger.error(f"Error al aprobar revisión de pañol: {e}")
        raise HTTPException(status_code=500, detail="Error al guardar la aprobación de pañol.")

    return RedirectResponse(url="/admin?msg=checklist_aprobado", status_code=status.HTTP_303_SEE_OTHER)


# ==========================================
# RUTAS DEL PANEL DE ADMINISTRACIÓN
# ==========================================

@app.get("/admin", response_class=HTMLResponse, tags=["Panel de Administración"])
async def admin_panel(
    request: Request,
    rut_docente: Optional[str] = Query(None),
    estado: Optional[str] = Query(None),
    taller_id: Optional[str] = Query(None),
    db: Session = Depends(get_db)
):
    """Panel de administración con métricas KPI y filtros."""
    query = db.query(models.Checklist)

    if rut_docente and rut_docente.strip():
        query = query.filter(models.Checklist.rut_docente == rut_docente.strip())
    if estado and estado.strip():
        query = query.filter(models.Checklist.estado == estado.strip())
    if taller_id and taller_id.strip():
        query = query.filter(models.Checklist.taller == taller_id.strip())

    checklists = query.order_by(models.Checklist.fecha_creacion.desc()).all()
    docentes = db.query(models.Docente).order_by(models.Docente.nombre).all()
    plantillas = obtener_plantillas(db)

    total_registros = len(checklists)
    total_finalizados = sum(1 for c in checklists if c.estado == "FINALIZADO")
    total_pendientes = sum(1 for c in checklists if c.estado == "PENDIENTE_REVISION")
    total_borradores = sum(1 for c in checklists if c.estado == "BORRADOR")

    return templates.TemplateResponse(
        request=request,
        name="admin.html",
        context={
            "checklists": checklists,
            "docentes": docentes,
            "plantillas": plantillas,
            "filtro_docente": rut_docente,
            "filtro_estado": estado,
            "filtro_taller": taller_id,
            "kpi": {
                "total": total_registros,
                "finalizados": total_finalizados,
                "pendientes": total_pendientes,
                "borradores": total_borradores,
            },
        },
    )


@app.post("/admin/eliminar/{checklist_id}", tags=["Panel de Administración"])
async def eliminar_checklist(checklist_id: int, db: Session = Depends(get_db)):
    """Elimina un checklist de la base de datos."""
    checklist = db.query(models.Checklist).filter(models.Checklist.id == checklist_id).first()
    if not checklist:
        raise HTTPException(status_code=404, detail="Checklist no encontrado")

    try:
        db.delete(checklist)
        db.commit()
    except SQLAlchemyError as e:
        db.rollback()
        logger.error(f"Error al eliminar checklist #{checklist_id}: {e}")
        raise HTTPException(status_code=500, detail="Error de base de datos al eliminar el checklist.")

    return RedirectResponse(url="/admin?msg=checklist_eliminado", status_code=status.HTTP_303_SEE_OTHER)


@app.get("/admin/export/excel", tags=["Panel de Administración"])
async def export_excel(
    rut_docente: Optional[str] = Query(None),
    estado: Optional[str] = Query(None),
    taller_id: Optional[str] = Query(None),
    db: Session = Depends(get_db)
):
    """Genera un reporte detallado en Excel (.xlsx) con auto-ajuste y formato profesional."""
    query = db.query(models.Checklist)

    if rut_docente and rut_docente.strip():
        query = query.filter(models.Checklist.rut_docente == rut_docente.strip())
    if estado and estado.strip():
        query = query.filter(models.Checklist.estado == estado.strip())
    if taller_id and taller_id.strip():
        query = query.filter(models.Checklist.taller == taller_id.strip())

    checklists = query.order_by(models.Checklist.fecha_creacion.desc()).all()

    wb = Workbook()
    ws = wb.active
    ws.title = "Reporte Aseo Gastronomía"

    headers = [
        "ID", "Taller", "RUT Docente", "Sección", 
        "Encargado Taller", "Fecha", "Estado", 
        "Pañolero a Cargo", "Obs. Pañol", "Fecha Revisión Pañol", "Fecha Creación"
    ]
    ws.append(headers)

    # Estilos de cabecera
    header_fill = PatternFill(start_color="0D6EFD", end_color="0D6EFD", fill_type="solid")
    header_font = Font(color="FFFFFF", bold=True)
    for cell in ws[1]:
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center")

    for c in checklists:
        f_rev = getattr(c, "fecha_revision_panolero", None)
        f_creac = getattr(c, "fecha_creacion", None)
        ws.append([
            c.id,
            c.taller,
            c.rut_docente,
            c.codigo_seccion,
            getattr(c, "encargado_taller", None) or "N/A",
            str(c.fecha) if c.fecha else "N/A",
            c.estado,
            getattr(c, "nombre_panolero", None) or "N/A",
            getattr(c, "observacion_panolero", None) or "N/A",
            f_rev.strftime("%Y-%m-%d %H:%M") if f_rev else "N/A",
            f_creac.strftime("%Y-%m-%d %H:%M") if f_creac else "N/A"
        ])

    # Auto-ajuste de ancho de columnas
    for col in ws.columns:
        max_len = max(len(str(cell.value or '')) for cell in col)
        col_letter = col[0].column_letter
        ws.column_dimensions[col_letter].width = max(max_len + 4, 12)

    stream = io.BytesIO()
    wb.save(stream)
    stream.seek(0)

    filename = f"reporte_aseo_{datetime.now().strftime('%Y%m%d_%H%M')}.xlsx"
    return StreamingResponse(
        stream,
        headers={"Content-Disposition": f'attachment; filename="{filename}"'},
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    )


@app.get("/admin/export/pdf/{checklist_id}", tags=["Panel de Administración"])
async def export_pdf(checklist_id: int, db: Session = Depends(get_db)):
    """Genera un informe PDF estilizado contemplando el doble checklist."""
    c = db.query(models.Checklist).filter(models.Checklist.id == checklist_id).first()
    if not c:
        raise HTTPException(status_code=404, detail="Checklist no encontrado")

    buffer = io.BytesIO()
    doc = SimpleDocTemplate(
        buffer, pagesize=letter, rightMargin=36, leftMargin=36, topMargin=36, bottomMargin=36
    )
    styles = getSampleStyleSheet()
    story = []

    title_style = ParagraphStyle(
        "CustomTitle",
        parent=styles["Heading1"],
        fontSize=16,
        textColor=colors.HexColor("#0d6efd"),
        spaceAfter=12,
    )
    cell_style = ParagraphStyle(
        "CellText",
        parent=styles["Normal"],
        fontSize=9,
        leading=11
    )

    story.append(Paragraph(f"Informe Auditado de Checklist - #{c.id}", title_style))
    story.append(Spacer(1, 10))

    f_rev = getattr(c, "fecha_revision_panolero", None)
    info_data = [
        [
            Paragraph("<b>Taller:</b>", cell_style), Paragraph(str(c.taller), cell_style),
            Paragraph("<b>Estado:</b>", cell_style), Paragraph(str(c.estado), cell_style)
        ],
        [
            Paragraph("<b>Docente (RUT):</b>", cell_style), Paragraph(str(c.rut_docente), cell_style),
            Paragraph("<b>Sección:</b>", cell_style), Paragraph(str(c.codigo_seccion), cell_style)
        ],
        [
            Paragraph("<b>Encargado Taller:</b>", cell_style), Paragraph(str(getattr(c, "encargado_taller", "") or "N/A"), cell_style),
            Paragraph("<b>Fecha Checklist:</b>", cell_style), Paragraph(str(c.fecha or "N/A"), cell_style)
        ],
        [
            Paragraph("<b>Pañolero A Cargo:</b>", cell_style), Paragraph(str(getattr(c, "nombre_panolero", "") or "Pendiente"), cell_style),
            Paragraph("<b>Fecha Revisión Pañol:</b>", cell_style), Paragraph(f_rev.strftime("%Y-%m-%d %H:%M") if f_rev else "Pendiente", cell_style)
        ],
        [
            Paragraph("<b>Obs. General Pañol:</b>", cell_style), Paragraph(str(getattr(c, "observacion_panolero", "") or "Sin observaciones"), cell_style),
            Paragraph("", cell_style), Paragraph("", cell_style)
        ]
    ]
    t_info = Table(info_data, colWidths=[110, 150, 110, 150])
    t_info.setStyle(TableStyle([
        ("BACKGROUND", (0, 0), (-1, -1), colors.HexColor("#f8f9fa")),
        ("GRID", (0, 0), (-1, -1), 0.5, colors.HexColor("#dee2e6")),
        ("PADDING", (0, 0), (-1, -1), 6),
    ]))
    story.append(t_info)
    story.append(Spacer(1, 15))

    story.append(Paragraph("<b>Detalle de Actividades Evaluadas (Doble Checklist):</b>", styles["Heading2"]))
    story.append(Spacer(1, 8))

    header_style = ParagraphStyle("HeaderStyle", parent=cell_style, textColor=colors.white, fontName="Helvetica-Bold")
    headers_act = [
        Paragraph("Actividad", header_style),
        Paragraph("Alumno Responsable", header_style),
        Paragraph("Estado Docente", header_style),
        Paragraph("Estado Pañol", header_style),
        Paragraph("Obs. Pañol", header_style)
    ]
    act_rows = [headers_act]

    actividades_data = safe_json_loads(getattr(c, "actividades_json", None))

    if actividades_data:
        for item in actividades_data:
            act_rows.append([
                Paragraph(str(item.get("actividad", "N/A")), cell_style),
                Paragraph(str(item.get("alumno_encargado", "N/A") or "N/A"), cell_style),
                Paragraph(str(item.get("estado", "N/A")), cell_style),
                Paragraph(str(item.get("estado_panol", "Pendiente")), cell_style),
                Paragraph(str(item.get("obs_panol", "-") or "-"), cell_style)
            ])
    else:
        act_rows.append([
            Paragraph("Sin detalle de actividades registradas", cell_style),
            Paragraph("-", cell_style),
            Paragraph("-", cell_style),
            Paragraph("-", cell_style),
            Paragraph("-", cell_style)
        ])

    t_act = Table(act_rows, colWidths=[140, 110, 85, 85, 100])
    t_act.setStyle(TableStyle([
        ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#0d6efd")),
        ("GRID", (0, 0), (-1, -1), 0.5, colors.HexColor("#cccccc")),
        ("PADDING", (0, 0), (-1, -1), 5),
    ]))
    story.append(t_act)

    doc.build(story)
    buffer.seek(0)

    filename = f"Checklist_{c.id}.pdf"
    return StreamingResponse(
        buffer,
        headers={"Content-Disposition": f'attachment; filename="{filename}"'},
        media_type="application/pdf",
    )


# ==========================================
# ENDPOINTS API REST (SELECTORES DINÁMICOS)
# ==========================================

@app.get("/api/docentes", response_model=List[DocenteOut], tags=["API Selectores"])
def obtener_docentes(db: Session = Depends(get_db)):
    """Retorna el listado de docentes ordenados por nombre."""
    return db.query(models.Docente).order_by(models.Docente.nombre).all()


@app.get("/api/docentes/{rut_docente}/secciones", response_model=List[SeccionOut], tags=["API Selectores"])
def obtener_secciones_por_docente(rut_docente: str, db: Session = Depends(get_db)):
    """Obtiene las secciones asignadas a un docente según su RUT."""
    secciones = (
        db.query(models.Seccion)
        .filter(models.Seccion.rut_docente == rut_docente)
        .order_by(models.Seccion.nombre_asignatura)
        .all()
    )
    return [
        {
            "cod_asignatura": s.cod_asignatura or "",
            "asignatura": s.nombre_asignatura,
            "seccion": s.codigo_seccion,
        }
        for s in secciones
    ]


@app.get("/api/secciones/{seccion_id}/alumnos", response_model=List[AlumnoOut], tags=["API Selectores"])
def obtener_alumnos_por_seccion(seccion_id: str, db: Session = Depends(get_db)):
    """Obtiene el listado de alumnos inscritos en una sección dada."""
    alumnos = (
        db.query(models.AlumnoInscrito)
        .filter(models.AlumnoInscrito.seccion == seccion_id)
        .order_by(models.AlumnoInscrito.nombre_alumno)
        .all()
    )
    return [{"rut": a.rut_alumno, "nombre": a.nombre_alumno} for a in alumnos]
